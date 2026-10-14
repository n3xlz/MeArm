# -*- coding: utf-8 -*-
"""把 markdown（受控子集）渲染为带原生 Word 公式对象的 .docx"""
import re
import sys
import os

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tex2omml import tex_to_omml

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import parse_xml, OxmlElement

SONG = "宋体"
HEI = "黑体"
WEST = "Times New Roman"

# ============================================================
#  行内文本解析：普通文字 / 粗体 / 数学 / 代码
# ============================================================
def inline_tokens(text, bold=False):
    out, buf, i = [], "", 0
    while i < len(text):
        c = text[i]
        if c == "$":
            j = text.find("$", i + 1)
            if j == -1:
                buf += c
                i += 1
                continue
            if buf:
                out.append(("t", buf, bold))
                buf = ""
            out.append(("m", text[i + 1:j], bold))
            i = j + 1
            continue
        if text.startswith("**", i):
            j = text.find("**", i + 2)
            if j == -1:
                buf += c
                i += 1
                continue
            if buf:
                out.append(("t", buf, bold))
                buf = ""
            out.extend(inline_tokens(text[i + 2:j], bold=True))
            i = j + 2
            continue
        if c == "`":
            j = text.find("`", i + 1)
            if j == -1:
                buf += c
                i += 1
                continue
            if buf:
                out.append(("t", buf, bold))
                buf = ""
            out.append(("c", text[i + 1:j], bold))
            i = j + 1
            continue
        buf += c
        i += 1
    if buf:
        out.append(("t", buf, bold))
    return out


def plain(text):
    return "".join(t[1] for t in inline_tokens(text) if t[0] != "m")


# ============================================================
#  样式
# ============================================================
def set_rfonts(el, ascii_font, ea_font):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    rf.set(qn("w:ascii"), ascii_font)
    rf.set(qn("w:hAnsi"), ascii_font)
    rf.set(qn("w:eastAsia"), ea_font)
    rf.set(qn("w:cs"), ascii_font)


def style_font(style, ascii_font, ea_font, size=None, bold=None, color=None):
    set_rfonts(style.element, ascii_font, ea_font)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


def build_styles(doc):
    st = doc.styles

    n = st["Normal"]
    style_font(n, WEST, SONG, size=12, bold=False, color=RGBColor(0, 0, 0))
    n.paragraph_format.line_spacing = 1.45
    n.paragraph_format.space_after = Pt(3)
    n.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, sb, sa in (("Heading 1", 16, 16, 10),
                               ("Heading 2", 14, 12, 8),
                               ("Heading 3", 12.5, 10, 6)):
        s = st[name]
        style_font(s, WEST, HEI, size=size, bold=True, color=RGBColor(0, 0, 0))
        s.paragraph_format.space_before = Pt(sb)
        s.paragraph_format.space_after = Pt(sa)
        s.paragraph_format.line_spacing = 1.2
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    cap = st.add_style("FigCaption", 1)  # 1 = paragraph
    cap.base_style = st["Normal"]
    style_font(cap, WEST, SONG, size=10.5, bold=False, color=RGBColor(0, 0, 0))
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.line_spacing = 1.15

    tb = st["Table Grid"]
    style_font(tb, WEST, SONG, size=10.5, bold=False, color=RGBColor(0, 0, 0))


# ============================================================
#  内容块渲染
# ============================================================
class Builder:
    def __init__(self, doc):
        self.doc = doc
        self.eq_no = 0

    # ---------- 段落 ----------
    def _fill(self, p, text, default_font=None):
        for kind, val, bold in inline_tokens(text):
            if kind == "m":
                p._p.append(parse_xml(tex_to_omml(val)))
            else:
                r = p.add_run(val)
                if bold:
                    r.font.bold = True
                if kind == "c":
                    set_rfonts(r._element, "Consolas", "Consolas")
                    r.font.size = Pt(10.5)
                elif default_font:
                    set_rfonts(r._element, WEST, default_font)
        return p

    def para(self, text, style=None, indent=None, align=None):
        p = self.doc.add_paragraph(style=style)
        self._fill(p, text)
        if indent:
            p.paragraph_format.first_line_indent = Cm(indent)
        if align:
            p.alignment = align
        return p

    def heading(self, text, level):
        p = self.doc.add_paragraph(style="Heading %d" % level)
        self._fill(p, text)
        return p

    def equation(self, tex, number=True):
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.tab_stops.add_tab_stop(Cm(7.8), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
        pf.space_before = Pt(6)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.25
        p.add_run("\t")
        p._p.append(parse_xml(tex_to_omml(tex)))
        if number:
            self.eq_no += 1
            p.add_run("\t(%d)" % self.eq_no)
        return p

    def quote(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.right_indent = Cm(0.3)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.3
        self._fill(p, text, default_font=SONG)
        for r in p.runs:
            r.font.size = Pt(10.5)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F5F9")
        pPr.append(shd)
        pPr.append(self._pborder())
        return p

    @staticmethod
    def _pborder():
        pbdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), "1F4E79")
        pbdr.append(left)
        return pbdr

    def figure(self, path, caption, width_cm):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.add_run().add_picture(path, width=Cm(width_cm))
        cp = self.doc.add_paragraph(style="FigCaption")
        self._fill(cp, caption)
        for r in cp.runs:
            r.font.size = Pt(10.5)

    def table(self, raw_rows):
        header = [c.strip() for c in raw_rows[0].strip("|").split("|")]
        body = []
        for row in raw_rows[2:]:
            cells = [c.strip() for c in row.strip("|").split("|")]
            if len(cells) < len(header):
                cells += [""] * (len(header) - len(cells))
            body.append(cells[:len(header)])

        t = self.doc.add_table(rows=1 + len(body), cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True

        for j, txt in enumerate(header):
            cell = t.cell(0, j)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(1)
            for kind, val, bold in inline_tokens(txt):
                if kind == "m":
                    p._p.append(parse_xml(tex_to_omml(val)))
                else:
                    r = p.add_run(val)
                    r.font.bold = True
                    r.font.size = Pt(10.5)
                    set_rfonts(r._element, WEST, HEI)
            self._shade(cell, "E8EEF6")

        for i, row in enumerate(body, start=1):
            for j, txt in enumerate(row):
                cell = t.cell(i, j)
                p = cell.paragraphs[0]
                p.alignment = (WD_ALIGN_PARAGRAPH.CENTER
                               if len(plain(txt)) <= 16 else WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(1)
                for kind, val, bold in inline_tokens(txt):
                    if kind == "m":
                        p._p.append(parse_xml(tex_to_omml(val)))
                    else:
                        r = p.add_run(val)
                        r.font.size = Pt(10.5)
                        r.font.bold = bold
                        set_rfonts(r._element, WEST, SONG)
        # 表头跨页重复 + 所有行禁止跨页拆分。
        # 不加这两条时，长表格中途换页会把表头行切成两半：
        # 上一页留半行文字，下一页开头孤零零剩个 "/mm"，很难看。
        for idx, row_obj in enumerate(t.rows):
            trPr = row_obj._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if idx == 0:
                trPr.append(OxmlElement("w:tblHeader"))

        # 表格前后留白
        self.doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return t

    @staticmethod
    def _shade(cell, color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), color)
        tcPr.append(shd)


# ============================================================
#  markdown 解析
# ============================================================
def parse_md(path):
    lines = open(path, encoding="utf-8").read().split("\n")
    blocks, i, n = [], 0, len(lines)
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", s)
            blocks.append(("fig", m.group(2), m.group(1)))
            i += 1
            continue
        if s.startswith("|"):
            buf = []
            while i < n and lines[i].strip().startswith("|"):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("table", buf))
            continue
        if s.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            blocks.append(("quote", " ".join(buf)))
            continue
        if s.startswith("#### "):
            blocks.append(("h3", s[5:]))
            i += 1
            continue
        if s.startswith("### "):
            blocks.append(("h2", s[4:]))
            i += 1
            continue
        if s.startswith("## "):
            blocks.append(("h1", s[3:]))
            i += 1
            continue
        if s.startswith("# "):
            blocks.append(("h1", s[2:]))
            i += 1
            continue
        if s.startswith("$$"):
            buf = s
            while not (buf.rstrip().endswith("$$") and len(buf.strip()) > 4):
                i += 1
                buf += " " + lines[i].strip()
            inner = buf.strip()[2:-2]
            blocks.append(("eq", inner.strip()))
            i += 1
            continue
        if s in ("---", "***"):
            i += 1
            continue
        blocks.append(("p", s))
        i += 1
    return blocks


# ============================================================
#  封面 / 页脚
# ============================================================
def add_footer_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.font.size = Pt(10.5)
    set_rfonts(r._element, WEST, SONG)
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = " PAGE "
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r._r.append(b)
    r._r.append(t)
    r._r.append(e)


def build_cover(doc):
    def line(text, size, font, bold=True, space_after=6, space_before=0):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        set_rfonts(r._element, WEST, font)
        return p

    for _ in range(4):
        line("", 12, SONG, bold=False, space_after=0)
    line("广东工业大学 · 粤港机器人联合学院", 14, HEI, space_after=4)
    line("第十二期二面考核（算法 / 控制 / 视觉方向）", 12, SONG, bold=False, space_after=40)
    line("任务四：机械臂运动学方程推导", 26, HEI, space_after=10)
    line("—— meArm 机械臂运动学正解与逆解", 14, HEI, space_after=60)
    for label, val in (("姓　　名", "钟林哲"), ("学　　号", "3126000947"),
                       ("专　　业", "自动化"), ("方向班级", "自动化类 5 班"),
                       ("日　　期", "2026 年 10 月 12 日")):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run("%s：%s" % (label, val))
        r.font.size = Pt(13)
        r.font.underline = True
        set_rfonts(r._element, WEST, SONG)
    doc.add_page_break()


# ============================================================
#  主流程
# ============================================================
def main():
    # 输入 markdown 与输出 docx 都从命令行参数取，图片目录跟 markdown 同级。
    # 这样脚本可以从任何地方调用，不再写死 /tmp/mission4。
    src = sys.argv[1] if len(sys.argv) > 1 else "task4.md"
    out = sys.argv[2] if len(sys.argv) > 2 else "任务四-运动学推导.docx"
    figdir = os.path.dirname(os.path.abspath(src))

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)
    sec.different_first_page_header_footer = True

    build_styles(doc)
    add_footer_page_number(sec)
    build_cover(doc)

    b = Builder(doc)
    first_h1 = True
    for blk in parse_md(src):
        kind = blk[0]
        if kind == "h1":
            if not first_h1:
                pass
            first_h1 = False
            b.heading(blk[1], 1)
        elif kind == "h2":
            b.heading(blk[1], 2)
        elif kind == "h3":
            b.heading(blk[1], 3)
        elif kind == "p":
            b.para(blk[1], indent=0.85)
        elif kind == "eq":
            b.equation(blk[1])
        elif kind == "quote":
            b.quote(blk[1])
        elif kind == "fig":
            path = os.path.join(figdir, blk[1])
            width = 15.5 if "fig1" in blk[1] else 14.2
            b.figure(path, blk[2], width)
        elif kind == "table":
            b.table(blk[1])

    doc.save(out)
    print("已生成:", out, "| 公式编号总数:", b.eq_no)


if __name__ == "__main__":
    main()
